#!/usr/bin/env python3
"""Verify generated manuals contain their declared visual evidence."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

REPO = Path(__file__).resolve().parents[2]
MANUALS = REPO / "docs/manuals/generated"
DIAGRAMS = REPO / "docs/manuals/assets/diagrams"
EXPECTED_DOCUMENTS = {
    "iRespond_Product_Documentation.docx",
    "iRespond_Technical_Documentation.docx",
    "iRespond_User_Manual_All_Roles.docx",
    "iRespond_Training_Manual_All_Roles.docx",
}
CONTRIBUTION_DOCUMENTS = EXPECTED_DOCUMENTS - {"iRespond_Technical_Documentation.docx"}


def document_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + table_cells)


def main() -> int:
    diagrams = sorted(path.name for path in DIAGRAMS.glob("*.png"))
    if not diagrams:
        print("No rendered diagrams found", file=sys.stderr)
        return 1

    errors = []
    for name in sorted(EXPECTED_DOCUMENTS):
        path = MANUALS / name
        if not path.is_file():
            errors.append(f"missing generated manual: {name}")
            continue
        text = document_text(path)
        for diagram in diagrams:
            if diagram not in text:
                errors.append(f"{name} does not embed {diagram}")
        if name in CONTRIBUTION_DOCUMENTS and "07-my-contribution-offers.png" not in text:
            errors.append(f"{name} omits the contribution-offers interface asset")

    if errors:
        print("\n".join(errors), file=sys.stderr)
        return 1
    print(f"Verified {len(EXPECTED_DOCUMENTS)} manuals with {len(diagrams)} rendered diagrams")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
